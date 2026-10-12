"""Réinitialise la base et importe le référentiel depuis l'Excel Sage.

Usage : python setup_db.py "<chemin vers BASE SAGE ADMIRAL.xlsx>"

ATTENTION : supprime TOUTES les tables (produits, gammes, OF, scans) avant l'import.
"""
import sqlite3
import os
import sys

from admiral import config
from admiral.db import schema
from admiral.db.depot import Depot
from admiral.domain import referentiel

DB_FILE = config.chemin_base(config.lire_config())

def setup_database(excel_file):
    import openpyxl

    print("Initialisation de la nouvelle base de données MES...")
    
    # Connect to the database (creates it if it doesn't exist)
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    
    # 1. Drop existing tables to start fresh
    tables = ['Produits', 'Operations', 'Postes', 'Gammes', 'CdeDetail', 'Scans', 'OF_Gammes']
    for table in tables:
        cursor.execute(f"DROP TABLE IF EXISTS {table}")
        
    # 2. Recréer le schéma de l'application (sans les postes par défaut : ils viennent de l'Excel)
    conn.execute("PRAGMA user_version = 0")
    schema.initialiser(conn, postes_par_defaut=False)
    
    print("Schéma créé avec succès.")

    # 3. Import Data from Excel
    if os.path.exists(excel_file):
        print(f"Lecture du fichier Excel : {excel_file}")
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        
        # Référentiel : mêmes règles de lecture que l'import depuis l'application (Administration)
        compte = Depot(conn).importer_referentiel(referentiel.analyser(referentiel.lire_classeur(excel_file)))
        print("Référentiel importé :", ", ".join(f"{n} {quoi}" for quoi, n in compte.items()))

        # Import CDEDETAIL (Initial OF creation)
        if 'CDEDETAIL' in wb.sheetnames:
            ws = wb['CDEDETAIL']
            of_counter = 1000
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0 or not row[0]: continue
                cde = str(row[0]).strip()
                prod = str(row[1]).strip() if row[1] else ""
                design = str(row[2]).strip() if row[2] else ""
                observ = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                
                id_of = f"OF-{of_counter}"
                of_counter += 1
                qte = 1 # Par défaut, car pas dans le Excel d'origine
                
                cursor.execute("INSERT INTO CdeDetail (ID_OF, CDE, PROD, DESIGN, OBSERV, Qte) VALUES (?, ?, ?, ?, ?, ?)", (id_of, cde, prod, design, observ, qte))
                
                # Générer la gamme pour cet OF
                cursor.execute("SELECT CodeOp, Ordre FROM Gammes WHERE RefProd = ? ORDER BY Ordre", (prod,))
                gamme_standard = cursor.fetchall()
                for step in gamme_standard:
                    cursor.execute("INSERT INTO OF_Gammes (ID_OF, Ordre, CodeOp) VALUES (?, ?, ?)", (id_of, step[1], step[0]))
                    
            print("Commandes Détails (OF) et Gammes OF importées.")

    conn.commit()
    conn.close()
    print("Base de données mise à jour avec succès.")

if __name__ == "__main__":
    if len(sys.argv) != 2 or not os.path.exists(sys.argv[1]):
        print(__doc__)
        sys.exit(1)
    print(f"Base cible : {DB_FILE}")
    if input("Toutes les données (OF, scans, gammes) seront supprimées. Tapez OUI pour continuer : ").strip() != "OUI":
        print("Annulé.")
        sys.exit(0)
    setup_database(sys.argv[1])
