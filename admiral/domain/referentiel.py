"""Lecture du référentiel Sage exporté en Excel (feuilles PRODUIT, OPERATION, POSTE, GAMPROD).

Les règles de lecture sont celles de l'import historique (setup_db.py) :
- chaque feuille a une ligne d'en-tête, les lignes sans valeur en colonne A sont ignorées ;
- POSTE n'a pas de colonne opération : elle se déduit du code (DEC1 -> DEC) ;
- GAMPROD : une ligne avec une référence en colonne A ouvre la gamme d'un produit, les lignes
  suivantes portent un code opération en colonne C.
"""
from dataclasses import dataclass, field

FEUILLES = ("PRODUIT", "OPERATION", "POSTE", "GAMPROD")


@dataclass
class Referentiel:
    produits: list = field(default_factory=list)    # [(RefProd, DesProd)]
    operations: list = field(default_factory=list)  # [(RefOp, LibelOp)]
    postes: list = field(default_factory=list)      # [(CodPost, DesPost, RefOp)]
    gammes: dict = field(default_factory=dict)      # RefProd -> [RefOp dans l'ordre]

    @property
    def vide(self):
        return not (self.produits or self.operations or self.postes or self.gammes)


def _texte(valeur):
    return str(valeur).strip() if valeur is not None else ""


def _cellule(ligne, index):
    return _texte(ligne[index]) if len(ligne) > index else ""


def _lignes_de_donnees(lignes):
    """Lignes après l'en-tête, avec une valeur en colonne A."""
    return [ligne for i, ligne in enumerate(lignes) if i > 0 and ligne and _texte(ligne[0])]


def analyser(feuilles):
    """feuilles : {nom de feuille: [lignes (tuples de valeurs)]} -> Referentiel."""
    ref = Referentiel()
    ref.produits = [(_cellule(l, 0), _cellule(l, 1)) for l in _lignes_de_donnees(feuilles.get("PRODUIT", []))]
    ref.operations = [(_cellule(l, 0), _cellule(l, 1)) for l in _lignes_de_donnees(feuilles.get("OPERATION", []))]
    for ligne in _lignes_de_donnees(feuilles.get("POSTE", [])):
        code = _cellule(ligne, 0)
        ref.postes.append((code, _cellule(ligne, 1), "".join(c for c in code if not c.isdigit()).strip()))

    produit_courant = None
    for i, ligne in enumerate(feuilles.get("GAMPROD", [])):
        if i == 0 or not ligne:
            continue
        if _cellule(ligne, 0):
            produit_courant = _cellule(ligne, 0)
            ref.gammes.setdefault(produit_courant, [])
        elif _cellule(ligne, 2) and produit_courant:
            ref.gammes[produit_courant].append(_cellule(ligne, 2))
    return ref


def lire_classeur(chemin):
    """Lit les feuilles utiles d'un classeur .xlsx (nécessite openpyxl)."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("La lecture des fichiers Excel nécessite le module openpyxl "
                           "(pip install -r requirements.txt).") from exc
    classeur = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    try:
        return {nom: list(classeur[nom].iter_rows(values_only=True)) for nom in FEUILLES if nom in classeur.sheetnames}
    finally:
        classeur.close()
